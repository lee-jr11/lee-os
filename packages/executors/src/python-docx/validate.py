import base64
import hashlib
import json
import sys
import zipfile
from io import BytesIO

try:
    from docx import Document
except ImportError:
    sys.exit(79)


FONT = "Times New Roman"
HALF_POINTS = 24
FONT_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MAIN_CONTENT_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def fail():
    raise ValueError("DOCX validation failed")


def direct_format(run):
    properties = run._element.rPr
    if properties is None or properties.rFonts is None:
        fail()
    size = properties.find(FONT_NS + "sz")
    size_cs = properties.find(FONT_NS + "szCs")
    if size is None or size_cs is None:
        fail()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        if properties.rFonts.get(FONT_NS + attribute) != FONT:
            fail()
    if size.get(FONT_NS + "val") != str(HALF_POINTS) or size_cs.get(FONT_NS + "val") != str(HALF_POINTS):
        fail()


def expected_blocks(document):
    blocks = [document["title"]]
    for section in document["sections"]:
        blocks.append(section["heading"])
        blocks.extend(section["paragraphs"])
    return blocks


def main():
    request = json.load(sys.stdin)
    raw = base64.b64decode(request["docxBase64"], validate=True)
    if not zipfile.is_zipfile(BytesIO(raw)):
        fail()
    with zipfile.ZipFile(BytesIO(raw)) as package:
        if package.testzip() is not None:
            fail()
        names = set(package.namelist())
        if "[Content_Types].xml" not in names or "word/document.xml" not in names or "word/styles.xml" not in names:
            fail()
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            fail()
        types = package.read("[Content_Types].xml")
        if MAIN_CONTENT_TYPE not in types:
            fail()

    parsed = Document(BytesIO(raw))
    if parsed.tables or len(parsed.paragraphs) != len(expected_blocks(request["document"])):
        fail()
    actual = [paragraph.text for paragraph in parsed.paragraphs]
    if actual != expected_blocks(request["document"]):
        fail()
    for paragraph in parsed.paragraphs:
        for run in paragraph.runs:
            if run.text:
                direct_format(run)
    for part in parsed.part.package.iter_parts():
        for relationship in part.rels.values():
            if relationship.is_external:
                fail()

    print(json.dumps({
        "schemaVersion": 1,
        "validator": "leeos-docx-validator-1",
        "format": "docx",
        "readable": True,
        "contentMatches": True,
        "fontFamily": FONT,
        "bodyFontSizePt": 12,
        "paragraphCount": len(parsed.paragraphs),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "bytes": len(raw),
    }, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        sys.exit(1)
