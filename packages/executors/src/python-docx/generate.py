import json
import sys
from io import BytesIO

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:
    sys.exit(79)


FONT = "Times New Roman"
SIZE = Pt(12)


def set_font(font):
    font.name = FONT
    font.size = SIZE


def set_run(run, value):
    run.text = value
    set_font(run.font)
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attribute}", FONT)
    size = r_pr.get_or_add_sz()
    size.set(qn("w:val"), "24")
    set_complex_size(r_pr)


def set_complex_size(r_pr):
    size = r_pr.find(qn("w:szCs"))
    if size is None:
        size = OxmlElement("w:szCs")
        r_pr.append(size)
    size.set(qn("w:val"), "24")


def main():
    request = json.load(sys.stdin)
    document = Document()
    for style in document.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            set_font(style.font)
            r_pr = style._element.get_or_add_rPr()
            fonts = r_pr.get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attribute}", FONT)
            size = r_pr.get_or_add_sz()
            size.set(qn("w:val"), "24")
            set_complex_size(r_pr)

    title = document.add_paragraph(style="Title")
    set_run(title.add_run(), request["title"])
    for section in request["sections"]:
        heading = document.add_paragraph(style="Heading 1")
        set_run(heading.add_run(), section["heading"])
        for value in section["paragraphs"]:
            paragraph = document.add_paragraph(style="Normal")
            set_run(paragraph.add_run(), value)

    document.core_properties.title = request["title"]
    output = BytesIO()
    document.save(output)
    sys.stdout.buffer.write(output.getvalue())


if __name__ == "__main__":
    main()
